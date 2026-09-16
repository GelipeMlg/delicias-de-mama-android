from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'entregables' / 'Resumen_base_de_datos_Delicias_de_Mama.docx'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin = s.bottom_margin = Cm(1.9)
s.left_margin = s.right_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 25), ('Subtitle', 12), ('Heading 1', 16), ('Heading 2', 12)]:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(6)
doc.core_properties.title = 'Base de datos de Delicias de Mamá'
doc.core_properties.subject = 'Resumen de diseño e implementación inicial para Proyecto Integrado TI'
doc.core_properties.author = 'Equipo Delicias de Mamá'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text, level=1):
    doc.add_heading(text, level)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for c, txt in zip(t.rows[0].cells, headers):
        c.text = txt
    for row in rows:
        for c, txt in zip(t.add_row().cells, row):
            c.text = txt
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:' + edge)
        for key, value in [('val','single'),('sz','4'),('color','D9D9D9')]:
            e.set(qn('w:' + key), value)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trPr.append(OxmlElement('w:tblHeader'))
        for c, width in zip(row.cells, widths):
            c.width = Cm(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = c._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd')
            sh.set(qn('w:fill'), '243746' if i == 0 else ('F0F4F7' if i % 2 else 'FFFFFF'))
            props.append(sh)
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m = OxmlElement('w:' + side)
                m.set(qn('w:w'), '85')
                m.set(qn('w:type'), 'dxa')
                margins.append(m)
            props.append(margins)
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(2)
                for run in par.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255,255,255)
    p('')

doc.add_heading('Base de datos de Delicias de Mamá', 0)
p('Proyecto Integrado TI · Resumen de diseño e implementación inicial', 'Subtitle')
p('Versión 1.0 · 16 de septiembre de 2026')
p('Equipo: Felipe Alvarado, Matias Otero y Jafet Lavados\nCarrera: Analista Programador · INACAP')
h('1 Propósito y resultado')
p('Construimos la estructura inicial de la base de datos para organizar clientes, pedidos personalizados, disponibilidad, cotizaciones y movimientos financieros de la repostería. El objetivo es mantener una misma información para la futura aplicación Android de clientes y dueña.')
p('Se crearon diez tablas en PostgreSQL, dentro del proyecto Supabase Delicias de mamá. La ejecución y las pruebas de integridad se comprobaron en el editor SQL. Las tablas quedaron sin datos de negocio y con protección de acceso habilitada. Esta entrega establece la estructura; no representa todavía una aplicación operativa.')
h('2 Correspondencia con el primer informe')
table(['Requisito', 'Tablas que lo respaldan'], [
 ('RF-01 Reservas','clientes, productos, pedidos y detalle_pedido'),
 ('RF-02 Disponibilidad','disponibilidad y pedidos'),
 ('RF-03 Personalización','detalle_pedido y pedidos'),
 ('RF-04 Cotizaciones','cotizaciones'),
 ('RF-05 Ventas','ventas y pagos'),
 ('RF-06 Gastos','gastos'),
 ('RF-07 Persistencia','Las diez tablas, sus claves y relaciones')
], [4.5,12.5])
p('RNF-02 se aborda inicialmente mediante RLS y restricción de permisos. Respaldos, restauración y acceso por rol siguen pendientes. RNF-01, RNF-03 y RNF-04 deberán comprobarse cuando exista la app Android y un escenario de prueba definido.')

doc.add_page_break()
h('3 Tablas de clientes y pedidos')
p('Cada tabla usa una clave primaria UUID, excepto disponibilidad, cuya clave es la fecha. Una clave foránea vincula un registro con otro y evita referencias inexistentes. Los campos creado_en registran el momento de alta.')
table(['Tabla', 'Campos principales y función'], [
 ('clientes','id, usuario_id, nombre, telefono, email y direccion. usuario_id enlaza opcionalmente con auth.users; permite conservar clientes registrados manualmente.'),
 ('productos','id, nombre, descripcion, categoria, imagen_url, precio_base y activo. El precio puede quedar sin definir cuando el producto requiere cotización.'),
 ('disponibilidad','fecha, bloqueada, capacidad_pedidos y observacion_interna. La capacidad se ingresa explícitamente por día; no se asignó un cupo real al negocio.'),
 ('pedidos','id, cliente_id, fecha_entrega, hora_entrega, tipo_entrega, direccion_entrega, especificaciones y estado. Un despacho exige dirección.'),
 ('detalle_pedido','id, pedido_id, producto_id, descripcion, cantidad, precio_unitario, sabores, rellenos, tematica y restricciones_alimentarias. producto_id es opcional para solicitudes fuera del catálogo.')
], [3.8,13.2])
h('Relaciones principales', 2)
p('Un cliente puede tener muchos pedidos. Cada pedido pertenece a un cliente y a una fecha configurada en disponibilidad. Una fecha puede relacionarse con muchos pedidos. Cada detalle pertenece a un pedido y, opcionalmente, a un producto del catálogo.')
h('Estados iniciales del pedido', 2)
p('solicitado, cotizado, confirmado, en_preparacion, entregado, rechazado y cancelado. El estado inicial es solicitado. La base limita los valores permitidos; las transiciones autorizadas entre estados se implementarán al validar el flujo con la dueña.')
p('La existencia de una fecha y una capacidad no evita por sí sola la sobreventa. La confirmación deberá verificar el cupo en una operación transaccional del servidor, incluso cuando dos clientes soliciten simultáneamente la misma fecha.')

doc.add_page_break()
h('4 Cotizaciones y control financiero')
table(['Tabla', 'Campos principales y función'], [
 ('cotizaciones','id, pedido_id, version, monto, detalle, valida_hasta y estado. Conserva versiones y admite como máximo una cotización aceptada por pedido.'),
 ('ventas','id, pedido_id, cotizacion_id, fecha_venta, total y estado. Como máximo una venta por pedido. La cotización vinculada debe corresponder al mismo pedido.'),
 ('pagos','id, venta_id, monto, tipo, metodo, referencia y fecha_pago. Una venta puede tener varios cobros o devoluciones.'),
 ('gastos','id, concepto, categoria, monto, fecha_gasto y observaciones. Clasifica insumos, costos fijos, costos variables y otros gastos.'),
 ('historial_pedido','id, pedido_id, estado_anterior, estado_nuevo, actor_usuario_id y ocurrido_en. Un disparador registra automáticamente el alta y los cambios de estado del pedido.')
], [3.8,13.2])
h('Decisiones de integridad', 2)
p('Los importes usan numeric(12,2), evitando la representación aproximada de double. La unidad monetaria propuesta para esta versión es el peso chileno. Las cantidades deben ser positivas; precios y totales no negativos; pagos y gastos mayores que cero.')
p('La combinación pedido y versión identifica cada cotización. Su fecha de vigencia no puede ser anterior al alta. Las referencias entre entidades restringen la eliminación de registros relacionados, sin borrado automático en cascada de la información comercial.')
h('Ventas y cobros tienen significados distintos', 2)
p('Una venta registra el valor de la operación. Un pago registra dinero cobrado o devuelto. Esta separación permite representar abonos sin duplicar ventas. El resumen financiero deberá distinguir ventas, cobros netos y gastos, y definir el período de cálculo antes de llamarlo rentabilidad.')
p('Quedan por implementar las reglas que impidan cobros excesivos, devoluciones superiores a lo cobrado y diferencias no autorizadas entre cotización, detalles y total de venta. Los reportes se calcularán a partir de las transacciones; no se creó una tabla con totales duplicados.')

doc.add_page_break()
h('5 Protección y verificación')
p('Se habilitó Row Level Security en las diez tablas y se retiraron los permisos de lectura y escritura de los roles public, anon y authenticated sobre ellas. No hay políticas de acceso de la app todavía. El propietario del proyecto puede trabajar desde el panel; una app con clave publicable no puede consultar estas tablas en esta etapa.')
p('El historial usa un disparador sin elevación de privilegios. Registra el usuario autenticado cuando existe; una operación administrativa sin sesión de la app puede dejar actor_usuario_id vacío. No constituye una auditoría completa ni inmutable.')
table(['Verificación ejecutada', 'Resultado'], [
 ('Diez tablas con RLS habilitada','Correcto'),
 ('Sin permisos SELECT, INSERT, UPDATE o DELETE para anon y authenticated en las diez tablas','Correcto'),
 ('Flujo de cliente, pedido, detalle, cotización, venta y pago','Correcto'),
 ('Historial automático de alta y cambio de estado','Correcto'),
 ('Rechazo de cantidad cero, gasto negativo y despacho sin dirección','Correcto'),
 ('Rechazo de venta duplicada, segunda cotización aceptada y pago sin venta existente','Correcto'),
 ('Reversión de los datos ficticios de las pruebas','Correcto')
], [13.3,3.7])
h('6 Próxima etapa', 1)
p('Validar con la dueña los cupos, anticipación mínima, retiro o despacho, vigencia de cotizaciones y abonos. Después implementar roles y políticas de acceso, confirmación transaccional, respaldo y restauración, y conexión Android. Agregar pruebas con sesiones de clientes diferentes y medición del tiempo de respuesta.')
h('7 Evidencias y referencias', 1)
p('Fuente funcional: InformeTrabajoIngSoft(V2).docx, fechado el 1 de septiembre de 2026.\nMigración: supabase/migrations/202609160001_base_inicial.sql\nPruebas: supabase/tests/base_inicial.sql\nProyecto: ivxeabovtxetvvxgyywa. Ejecución manual en el editor SQL el 16 de septiembre de 2026; no registrada automáticamente en el historial de Supabase CLI.')
p('Referencia técnica: Supabase, Row Level Security.\nhttps://supabase.com/docs/guides/database/postgres/row-level-security\nConsulta: 16 de septiembre de 2026.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
