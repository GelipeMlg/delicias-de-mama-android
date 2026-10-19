from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]
REF=ROOT/'Archivos Proyecto Integrado/01 Pautas/Carta-1-vert.docx'
OUT=ROOT/'Archivos Proyecto Integrado/02 Informes Word/Informe_U2_Diseno_Delicias_de_Mama_2026.docx'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';N={'w':W}
def el(tag,**attrs):
 x=E.Element('{'+W+'}'+tag)
 for k,v in attrs.items():x.set('{'+W+'}'+k,str(v))
 return x
with ZipFile(REF) as z:parts={n:z.read(n) for n in z.namelist()}
tree=E.fromstring(parts['word/document.xml']);body=tree.find('w:body',N);src=list(body)
sect=deepcopy(src[-1]);cover=[deepcopy(x) for x in src[:9]];closing=deepcopy(src[66])
for x in list(body):body.remove(x)
def replace(p,text):
 run=p.find('w:r',N);rp=deepcopy(run.find('w:rPr',N)) if run is not None and run.find('w:rPr',N) is not None else None
 for x in list(p):
  if x.tag!='{'+W+'}pPr':p.remove(x)
 r=el('r')
 if rp is not None:r.append(rp)
 for i,s in enumerate(text.split('\n')):
  if i:r.append(el('br'))
  t=el('t');t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=s;r.append(t)
 p.append(r);return p
def prop(p):
 pp=p.find('w:pPr',N)
 if pp is None:pp=el('pPr');p.insert(0,pp)
 return pp
def setprop(pp,tag,**attrs):
 old=pp.find('w:'+tag,N)
 if old is not None:pp.remove(old)
 pp.append(el(tag,**attrs))
def para(text,kind='body',newpage=False):
 p=replace(deepcopy(src[24 if kind=='title' else 27 if kind=='sub' else 25]),text);pp=prop(p)
 setprop(pp,'spacing',before=160 if kind=='title' else 100 if kind=='sub' else 0,after=160 if kind=='title' else 100,line=260,lineRule='auto')
 if kind=='title':
  r=p.find('w:r',N);rp=r.find('w:rPr',N)
  if rp is None:rp=el('rPr');r.insert(0,rp)
  setprop(rp,'rFonts',ascii='Tahoma',hAnsi='Tahoma',cs='Tahoma')
 if kind!='body':setprop(pp,'keepNext')
 if newpage:setprop(pp,'pageBreakBefore')
 body.append(p);return p
def page(t):
 spacer=para('',newpage=True);setprop(prop(spacer),'spacing',before=0,after=0,line=240,lineRule='exact');setprop(prop(spacer),'keepNext')
 para(t,'title')
def sub(t):para(t,'sub')
def p(t):para(t)
def table(headers,rows,widths):
 tb=el('tbl');pr=el('tblPr');pr.append(el('tblW',w=sum(widths),type='dxa'));pr.append(el('tblLayout',type='fixed'));borders=el('tblBorders')
 for s in ['top','left','bottom','right','insideH','insideV']:borders.append(el(s,val='single',sz=4,color='D9D9D9'))
 pr.append(borders);margins=el('tblCellMar')
 for s in ['top','bottom']:margins.append(el(s,w=90,type='dxa'))
 for s in ['left','right']:margins.append(el(s,w=110,type='dxa'))
 pr.append(margins);tb.append(pr);grid=el('tblGrid')
 for w in widths:grid.append(el('gridCol',w=w))
 tb.append(grid)
 for i,row in enumerate([headers]+rows):
  tr=el('tr');trp=el('trPr');trp.append(el('cantSplit'))
  if i==0:trp.append(el('tblHeader'))
  tr.append(trp)
  for text,w in zip(row,widths):
   tc=el('tc');tcp=el('tcPr');tcp.append(el('tcW',w=w,type='dxa'))
   if i==0:tcp.append(el('shd',fill='EEEEEE'))
   tc.append(tcp);q=replace(deepcopy(src[25]),text);setprop(prop(q),'spacing',before=0,after=0,line=245,lineRule='auto')
   if i==0:
    for rp in q.findall('w:r/w:rPr',N):rp.append(el('b'))
   tc.append(q);tr.append(tc)
  tb.append(tr)
 body.append(tb);para('')
from docx import Document
from docx.shared import Inches
R='http://schemas.openxmlformats.org/package/2006/relationships'
rels=E.fromstring(parts['word/_rels/document.xml.rels'])
figs=ROOT/'Archivos Proyecto Integrado/03 Documentacion/U2/figuras'
fig_count=0
def figure(name,caption):
 global fig_count
 fig_count+=1
 temp=Document();q=temp.add_paragraph();q.add_run().add_picture(str(figs/(name+'.png')),width=Inches(7.25))
 node=deepcopy(q._p)
 for x in node.iter():
  if x.tag.endswith('}blip'):x.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed','rIdU2fig'+str(fig_count))
  if x.tag.endswith('}docPr'):x.set('id',str(500+fig_count));x.set('name',name);x.set('descr',caption)
 setprop(prop(node),'jc',val='center');body.append(node)
 rel=E.SubElement(rels,'{'+R+'}Relationship');rel.set('Id','rIdU2fig'+str(fig_count));rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/image');rel.set('Target','media/u2_'+name+'.png')
 parts['word/media/u2_'+name+'.png']=(figs/(name+'.png')).read_bytes()
 p(caption)
replace(cover[2],'DISEÑO DE LA SOLUCIÓN DE SOFTWARE')
replace(cover[3],'Delicias de Mamá · Unidad 2')
replace(cover[8],'NOMBRE: Felipe Alvarado y Jafet Lavados\nCARRERA: Analista Programador · INACAP Punta Arenas\nASIGNATURA: Proyecto Integrado TI · TIHV43\nPROFESORA: Jessica Paola Miranda Márquez\nFECHA: 13 de octubre de 2026 · Versión para revisión')
for x in cover:body.append(x)
page('Resumen y alcance del diseño')
p('Diseñamos una solución Android para coordinar el catálogo, las solicitudes personalizadas, la disponibilidad diaria, las cotizaciones y los registros financieros de Delicias de Mamá. Este informe desarrolla las seis actividades de la Unidad 2 y conecta sus decisiones con los requisitos del diagnóstico U1 [1–2]. La propuesta utiliza Kotlin y Jetpack Compose, Supabase y Firebase Cloud Messaging.')
p('La app ya cuenta con una implementación avanzada. Por ello, distinguimos la arquitectura reconstruida desde los archivos actuales de los ajustes y validaciones propuestos. Los diagramas explican la solución vigente; el cronograma organiza su consolidación y no presenta como futuras tareas que ya se realizaron. La revisión documental no certifica el estado de la base remota ni la aceptación comercial de la dueña.')
p('El modelo incremental permite cerrar conjuntos útiles: acceso y catálogo, pedidos y cotizaciones, agenda y finanzas, y comunicación entre dispositivos. Cada incremento exige requisitos claros, revisión de diseño y pruebas de regresión. El aporte tecnológico consiste en compartir información y controlar permisos y cupos, con una interfaz apropiada para repostería y coctelería salada.')
sub('Objetivo de la unidad')
p('Especificar un diseño trazable y verificable de la solución, mediante planificación, vistas de arquitectura, interfaces, modelo relacional, controles de seguridad y casos de prueba. Se prioriza el recorrido cliente–dueña y la coherencia de la información antes de ampliar el alcance.')
sub('Contenido')
p('1 Propuestas y priorización\n2 Metodología y cronograma\n3 Arquitectura 4+1\n4 Interfaces y navegación\n5 Modelo de datos e integración\n6 Programación segura\n7 Plan de pruebas\n8 Conclusiones y cobertura de la pauta\nReferencias y fuentes')
p('Las figuras 1 a 9 son elaboración del equipo sobre archivos revisados el 13 de octubre de 2026. El prototipo de la figura 7 es conceptual y no una captura de ejecución. No se incorporan credenciales, registros de clientes ni trabajos de otras asignaturas.')
page('1 Propuestas y priorización')
sub('1 1 Alternativas de solución')
p('La siguiente reconstrucción analítica organiza ideas relacionadas con las necesidades de U1. No constituye un acta de brainstorming realizado en clase ni una votación de la contraparte. Debe contrastarse con el equipo y la docente [1–2].')
table(['Categoría','Idea','Decisión de alcance'],[
['Coordinación','Agenda, cupos y pedidos compartidos','Núcleo: evita información aislada.'],
['Oferta','Catálogo editable, fotos y modo pastel','Núcleo: adapta la oferta a productos distintos.'],
['Comunicación','Cotización revisable y notificaciones','Núcleo: facilita respuesta y seguimiento.'],
['Administración','Registro de ventas y gastos','Núcleo: consulta operativa por período.'],
['Comercial','Pagos en línea, abonos y fidelización','Diferido: requiere reglas y conciliación.'],
['Distribución','Rutas y seguimiento de despacho','Diferido: no integra el flujo vigente.']],[2300,4050,4250])
sub('1 2 Matriz ponderada propuesta')
p('Escala 1–5: mayor puntaje indica mayor contribución. Impacto I = 40 %, factibilidad F = 35 % y urgencia U = 25 %. Total = 0,40 I + 0,35 F + 0,25 U. Los valores son juicio de diseño del equipo, no mediciones del negocio. Seguridad es condición transversal, no una opción que compite por puntaje.')
table(['Propuesta','I','F','U','Total'],[
['Pedidos y agenda compartida','5','4','5','4,65'],['Catálogo y personalización','5','5','4','4,75'],['Cotizaciones y avisos','5','4','4','4,40'],['Ventas y gastos','4','4','4','4,00'],['Pagos en línea','4','2','2','2,80'],['Rutas de despacho','2','2','1','1,75']],[5700,1200,1200,1200,1300])
p('Seleccionamos las cuatro primeras por relación directa con el problema y componentes ya disponibles. El orden técnico respeta dependencias: identidad y catálogo preceden a pedido; cotización precede a aceptación y entrega. Una puntuación alta no elimina esas dependencias.')
page('2 Metodología y cronograma')
p('Mantenemos desarrollo incremental: cada conjunto de funciones se diseña, construye, revisa e integra con lo anterior. Git conserva versiones; las pruebas de reglas y recorridos evitan que una mejora altere funciones existentes. El trabajo se controla por entregables verificables, sin atribuir ceremonias ni aprobaciones inexistentes.')
sub('2 1 Plan propuesto de consolidación de U2')
p('No hay fecha oficial de entrega. Proponemos trabajar del 14 al 28 de octubre de 2026 y reprogramar con la docente. Las 15 horas corresponden a la dedicación académica orientativa de la guía, distribuida en 6 + 6 + 3 horas; no estiman el desarrollo completo de la app ni horas-persona multiplicadas por tres.')
table(['ID y actividad','Fecha propuesta','Horas','Depende de'],[
['A Ideas y prioridades','14 oct','2','Diagnóstico U1'],['B Ajustar cronograma y alcance','15 oct','4','A'],['C Arquitectura y diagramas','19 oct','3','B'],['D Interfaces y prototipo','20 oct','3','B; revisión C'],['E Datos e integración','26 oct','1','C y D'],['F Controles y plan de pruebas','27 oct','1,5','E'],['G Revisión cruzada y cierre','28 oct','0,5','A–F']],[4700,2500,1200,2200])
sub('2 2 Carta Gantt de consolidación')
table(['Trabajo','14–15 oct','19–20 oct','26–28 oct'],[
['A + B: alcance y planificación','■■ H1','',''],['C + D: arquitectura e interfaces','','■■ H2',''],['E + F: datos y verificación','','','■■'],['G: cierre de diseño','','','◆ H3']],[4600,2000,2000,2000])
p('H1: alcance y prioridades revisados. H2: diagramas y prototipo coherentes entre sí. H3: seis criterios U2 trazables y pendientes identificados. La red de dependencias es A → B → (C, D) → E → F → G. C y D pueden solaparse con coordinación; un desacuerdo en reglas de negocio obliga a revisar E y F.')
page('2 Recursos y control de avances')
sub('2 3 Responsabilidades propuestas')
p('La siguiente asignación es una propuesta para acordar, no una declaración de aportes pasados. Todos revisan el documento y deben poder explicar el diseño completo.')
table(['Integrante','Responsabilidad propuesta','Revisión cruzada'],[
['Felipe Alvarado','Arquitectura Android, integración, componentes e interfaces.','Revisar casos de prueba con Jafet.'],['Jafet Lavados','Modelo de datos, reglas del servidor y registro de pruebas.','Revisar arquitectura con Felipe.']],[2700,4500,3400])
sub('2 4 Recursos y dependencias externas')
p('Se requiere PC con Android Studio, repositorio Git, archivos SQL, acceso autorizado a Supabase/Firebase y dos dispositivos Android para las pruebas posteriores. La red es necesaria para operaciones compartidas. La validación de la dueña requiere una sesión acordada y datos ficticios; no depende de adquirir nuevas funciones comerciales.')
table(['Riesgo de planificación','Respuesta','Criterio de cierre'],[
['Falta fecha oficial','Acordar calendario y conservar versión propuesta.','Fecha y sección registradas antes de entrega.'],['Reglas de cupo ambiguas','Confirmar si un pedido equivale a una unidad de capacidad.','Regla aprobada y reflejada en RF y prueba.'],['Entorno remoto no coincide con SQL','Inventariar versiones antes de migrar.','No aplicar migraciones a ciegas.'],['Falla de red o dispositivo','Preparar entorno de ensayo y registrar bloqueos.','No marcar aprobado un caso no ejecutado.']],[2900,4100,3600])
sub('2 5 Gestión de cambios')
p('Cada cambio registra solicitante, fecha, motivo, requisito, impacto y decisión. Primero se ajustan alcance y criterios; luego diagrama, código, migración y prueba. Las reuniones de revisión propuestas ocurren en H1 y H2, con cierre en H3. El historial Git permite comparar versiones, pero no sustituye actas de aceptación ni respaldo de datos de negocio.')
page('3 Arquitectura y vista lógica')
p('Aplicamos 4+1: vista lógica, de desarrollo, de procesos y física, articuladas por escenarios. La guía concentra el detalle en las tres primeras más el caso de uso del sistema; añadimos un bosquejo físico para explicar la nube. Las clases de la figura 1 corresponden al dominio Android [3].')
figure('01_clases','Figura 1. Diagrama de clases del dominio Android y asociaciones por identificador. Se omiten propiedades secundarias para legibilidad.')
p('ShopState reúne colecciones de Product, Order, DayConfig, Sale y Expense. ShopRules valida cambios locales; el servidor repite reglas críticas porque el cliente puede alterarse. Order conserva nombre de producto, opciones, estado y revisión. Su representación combina varias tablas mediante RPC; una clase de pantalla no equivale necesariamente a una tabla.')
p('La asociación Sale–Order es opcional para ventas manuales. Expense no expone orderId en Android, aunque el esquema relacional admite atribución opcional de gastos. Se conserva esta diferencia para no declarar una función de interfaz inexistente.')
page('3 Vista de desarrollo')
figure('02_componentes','Figura 2. Diagrama de componentes y dependencias principales. Flecha discontinua: implementación de interfaz; flechas continuas: uso.')
p('La pantalla emite acciones a ShopViewModel y observa UiState. El coordinador aplica reglas, ejecuta una operación y publica carga, resultado o error. ShopRepository define load/save; ConnectedRepository adapta el estado al contrato JSON de las funciones PostgreSQL. MainActivity selecciona este repositorio para la app conectada.')
p('SupabaseAuthRepository administra identidad y consulta el privilegio de dueña. SecureSessionStore protege la sesión opcional. ProductPhotos gestiona fotos y NotificationSupport los avisos. Esta separación facilita comprobar reglas sin depender de la interfaz, aunque las pruebas locales no sustituyen la integración con los servicios.')
sub('Decisión de diseño')
p('Mantener reglas críticas en transacciones del servidor protege frente a dos dispositivos que operan al mismo tiempo. El repositorio acepta un cambio por guardado, lo que simplifica identificar la operación que falló. La vista no debe anunciar éxito hasta que el servidor confirme.')
page('3 Vista de procesos')
figure('03_bpmn','Figura 3. BPMN simplificado del recorrido feliz y rechazo inicial, en un proceso con tres carriles. Círculo fino: inicio; borde grueso: fin; rectángulo redondeado: tarea; rombo X: decisión exclusiva.')
p('El cliente envía su solicitud; el sistema valida opciones y capacidad. Si falla, informa rechazo sin registrar un pedido válido. Si acepta, guarda y ocupa el cupo antes de la cotización. La dueña cotiza, el cliente revisa y el servidor valida la revisión antes de confirmar. Entregar registra la venta una sola vez.')
p('Excepciones: una cotización modificada obliga a revisar de nuevo; un pedido cerrado no admite recotización; cancelar libera ocupación conforme a las reglas. La figura omite estos subflujos para legibilidad, pero se incluyen en el plan de pruebas. La entrega actual no incorpora estados de preparación o listo para retiro.')
page('3 Vista de escenarios')
figure('04_casos','Figura 4. Caso de uso de contexto CU 0 y objetivos de los actores dentro del límite del sistema. La autenticación es una precondición de las operaciones privadas.')
sub('CU 0 Gestionar un encargo personalizado')
p('Actores principales: cliente y dueña. Precondiciones: identidad válida, producto activo y fecha disponible. Disparador: cliente solicita un producto. Flujo: consultar oferta → personalizar → solicitar → cotizar → revisar y aceptar → entregar. Resultado: pedido entregado y venta asociada sin duplicación.')
p('Alternativas: entrada inválida, cupo agotado, revisión obsoleta, cancelación o pérdida de conexión. El sistema conserva un estado válido y permite consultar nuevamente. Supabase y FCM son servicios de soporte del escenario; recibir un aviso no equivale a aceptar una cotización ni a registrar un pago.')
page('3 Vista física e infraestructura')
figure('09_despliegue','Figura 5. Despliegue lógico de infraestructura y canales de comunicación. No representa una auditoría de la configuración remota.')
p('Supabase provee servicios administrados de backend: Auth, PostgreSQL, Storage, Realtime y ejecución de funciones. Firebase Cloud Messaging transporta avisos. No desplegamos máquinas virtuales propias; por ello no atribuimos una capa IaaS administrada por el equipo. Es más preciso describir estos componentes como backend y plataforma administrados que forzar su clasificación como IaaS.')
p('El cliente conserva sesión protegida si se elige recordar acceso; los registros compartidos se consultan al servidor. Realtime indica cambios y provoca una nueva lectura; no transporta por sí mismo toda la lógica comercial. El envío push se origina en el servidor y puede llegar en segundo plano, sujeto a permisos y restricciones del dispositivo.')
p('Ante caída de conectividad se debe informar que la operación no está confirmada y recargar antes de reintentar. Respaldo y restauración de PostgreSQL, usuarios y objetos de Storage requieren procedimientos propios; un ZIP de GitHub no recupera esos datos.')
page('4 Interfaces y navegación')
figure('05_navegacion','Figura 6. Mapa de navegación según perfiles y pestañas de MainActivity.kt. Los formularios aparecen dentro de esos recorridos.')
p('El cliente necesita consultar oferta y disponibilidad, personalizar sin ambigüedad y revisar su cotización. La dueña requiere actualizar catálogo, cupos y movimientos y responder pedidos. El visitante consulta información pública; para enviar o gestionar pedidos debe iniciar sesión. La navegación cambia únicamente después de verificar identidad y privilegio.')
p('La jerarquía prioriza catálogo/productos y mantiene agenda y pedidos visibles. Finanzas solo aparece para la dueña. Un aviso conduce a Pedidos; volver a la app actualiza información. El gesto de deslizar hacia abajo permite recargar, y los estados de carga o error deben conservar una explicación comprensible.')
page('4 Prototipo y decisiones UX')
figure('06_prototipo','Figura 7. Prototipo conceptual de tres pantallas esenciales. Importes y opciones son ejemplos; no es evidencia de una sesión real ni de aprobación de la dueña.')
p('La identidad visual combina fondo crema #FFF8F4, texto café #38251F, rosa #984858 y tarjetas redondeadas. La propuesta retoma los colores presentes en MainActivity.kt. El botón principal usa un verbo concreto: personalizar, enviar o guardar. La foto del producto apoya la selección sin reemplazar su nombre y descripción.')
p('Modo de pastel habilita un sabor y un relleno seleccionables por cliente, a partir de listas configuradas por la dueña. Porciones se presenta una sola vez: bajo modo pastel o como opción general según corresponda. Se evita obligar a un producto salado a usar sabores de torta.')
p('Validación UX propuesta: observar a la dueña crear un producto y corregir una cotización, y a un cliente enviar un pedido. Registrar errores, ayuda y comprensión del estado. Comprobar texto ampliado, TalkBack, contraste y controles táctiles; estos resultados siguen pendientes y no se infieren del prototipo.')
page('5 Modelo relacional comercial')
p('Elegimos PostgreSQL porque pedidos, clientes, fechas, cotizaciones y movimientos requieren integridad referencial y transacciones. Las fotografías se almacenan como objetos en Storage y la base conserva su referencia; no se justifica una segunda base documental para este alcance [3].')
figure('07_datos','Figura 8. Submodelo comercial con claves y cardinalidades. PK: primaria; FK: foránea; UK: única. Flechas expresan relación entre entidades, no flujo de datos.')
p('Cada pedido tiene cliente y fecha. Puede tener varios detalles y versiones de cotización; el cliente Android actual trabaja con el primer detalle. Una venta puede ser manual o enlazar un pedido; la unicidad de pedido_id evita ventas duplicadas. La FK compuesta (cotizacion_id, pedido_id) exige coherencia entre venta y cotización. Por esa restricción y la unicidad, una cotización puede asociarse como máximo con una venta.')
p('Pagos existe en el esquema, pero no forma parte del flujo implementado de abonos/cobro en Android. El historial de estado se mantiene en historial_pedido y se detalla en el diccionario. La figura es un submodelo legible, no el inventario completo de tablas.')
page('5 Diccionario de datos comerciales')
table(['Entidad y clave','Atributos esenciales','Regla y uso'],[
['clientes · id UUID','usuario_id UUID único; nombre; teléfono; email; dirección.','usuario_id referencia auth.users y admite NULL. No guardar contraseñas aquí.'],
['productos · id UUID','nombre; descripción; precio_base numeric; activo; imagen_url; sabores_habilitados; sabores[]; rellenos[]; porciones[].','Precio no negativo. Opciones validadas por RPC; foto separada de sus bytes.'],
['disponibilidad · fecha date','bloqueada boolean; capacidad_pedidos integer; observacion_interna.','Un registro por fecha; cupo validado en operación transaccional.'],
['pedidos · id UUID','cliente_id; fecha_entrega; hora; tipo_entrega; estado; revision; nombre_contacto; contacto.','FK cliente y fecha; revisión detecta cambio concurrente.'],
['detalle_pedido · id UUID','pedido_id; producto_id; descripción; cantidad; sabores; rellenos; temática; restricciones; personas.','Cantidad positiva. Conserva descripción del producto solicitado.'],
['cotizaciones · id UUID','pedido_id; version; monto; detalle; estado; valida_hasta.','Única versión por pedido; como máximo una aceptada.'],
['ventas · id UUID','pedido_id nullable único; cotizacion_id; concepto; total; fecha_venta; estado.','Venta por pedido o manual. Anulación conserva registro.'],
['gastos · id UUID','concepto; categoría; monto; fecha_gasto; pedido_id nullable.','FK opcional impide vínculo inexistente; UI actual no la expone.']],[2700,4300,3600])
p('Los importes SQL usan numeric(12,2); Android utiliza Long y valida pesos enteros CLP. Esta restricción de interfaz debe mantenerse explícita para no perder decimales inesperados. El esquema admite más estados y tipo despacho que el recorrido simplificado actual; no se declaran funciones terminadas por existir una columna.')
page('5 Identidad y notificaciones')
figure('08_push','Figura 9. Submodelo de identidad y cola push. Cada aviso tiene un destinatario y pedido; cada envío vincula aviso y dispositivo, con par único.')
table(['Entidad complementaria','Atributos y relación'],[
['historial_pedido','id; pedido_id; estado_anterior/nuevo; actor_usuario_id; ocurrido_en. Trigger de alta/cambio de estado.'],
['pagos','id; venta_id; monto; tipo cobro/devolución; método; fecha_pago. Funcionalidad de app diferida.'],
['app_changes','id=1; version bigint. Señal técnica de invalidación; no copia registros comerciales.'],
['Storage','Objetos de fotos y referencia en productos. Gestión de permisos independiente del catálogo.']],[3300,7300])
p('envios_push_app conserva enviado, intentos, disponible_en y lease para gestionar el procesamiento. El token pertenece a un usuario y vence; se debe comprobar su registro y retiro al cambiar de sesión. La cola permite separar guardar un pedido de entregar un aviso; un fallo de FCM no debe deshacer una operación comercial confirmada.')
page('5 Integración y consistencia')
sub('5 1 Contratos de lectura y escritura')
table(['Contrato','Responsabilidad','Comprobación requerida'],[
['estado_app','Devuelve catálogo, pedidos permitidos, agenda, ocupación, finanzas autorizadas y avisos propios.','Dos clientes no recuperan pedidos del otro.'],
['fotos_catalogo_app','Obtiene referencias de fotografías del catálogo.','Solo contenido público previsto y productos autorizados.'],
['guardar_cambio_app','Aplica un cambio con estado anterior/siguiente para pedido, día o finanzas.','Permiso, revisión y reglas en servidor.'],
['guardar_producto_con_foto_app','Actualiza catálogo y referencia de foto.','Rol de dueña y validación de recursos.'],
['app_changes + Realtime','Indica nueva versión y provoca recarga.','Reabrir app recupera estado aunque se pierda un evento.']],[3100,4200,3300])
sub('5 2 Reglas transaccionales')
p('Solicitar debe validar producto, opciones, fecha y capacidad dentro del servidor. Se bloquea el recurso correspondiente para evitar comprometer el último cupo desde dos sesiones. La revisión vigente impide aceptar información que cambió durante la lectura. Entregar enlaza una venta de forma única. Los detalles de esta lógica se ubican en la migración 202609260002_pedidos_compartidos.sql [3].')
p('Una foto se transfiere a Storage antes de asociarla al producto. Base y Storage no forman una sola transacción: deben contemplarse fallo de carga, fallo de asociación y objetos huérfanos. Se recomienda revisar esos recursos mediante un procedimiento controlado; no borrar por similitud de nombre.')
sub('5 3 Evolución y límites')
p('El modelo se reconstruyó con siete migraciones locales, desde base_inicial hasta fotos_productos. Antes de aplicar cambios se compara el esquema remoto con el historial real de ejecución. Las migraciones son incrementales y se prueban en un entorno aislado. Ni este informe ni cambiar de rama Git modifican automáticamente la base de datos.')
page('6 Estándares de programación segura')
table(['Perfil','Lectura','Cambios permitidos'],[
['Visitante','Catálogo activo y disponibilidad pública.','Sin operaciones comerciales privadas.'],
['Cliente autenticado','Información pública y pedidos propios.','Solicitar, confirmar o cancelar según estado.'],
['Dueña autorizada','Pedidos y administración del negocio.','Catálogo, agenda, cotizar, entregar y finanzas.'],
['Servicio de notificaciones','Cola y dispositivos necesarios.','Procesar envíos; secreto solo en servidor.']],[2600,3600,4400])
p('La interfaz oculta funciones por perfil, pero el control decisivo ocurre en auth.uid(), es_duena() y las RPC. Las tablas tienen RLS y permisos restringidos. Las funciones SECURITY DEFINER deben revisar explícitamente al solicitante; activar RLS no basta para proteger una función privilegiada [3,5].')
table(['Estándar propuesto','Aplicación y verificación'],[
['Validar entradas','Longitudes, importes, fechas, listas permitidas y tipo JSON. Servidor revalida; no confiar en controles de pantalla.'],
['Proteger identidad','Auth verifica credenciales. Bcrypt es hash no reversible; sesiones recordadas usan AES-GCM y clave de Android Keystore.'],
['Privilegio mínimo','Cliente no se autoconcede rol. Probar RPC y acceso directo, no solo botones. Nunca incluir service_role ni secretos FCM en APK.'],
['Errores seguros','Mensaje útil sin SQL, tokens o datos ajenos. Conservar identificador técnico para diagnóstico, no contraseña ni contenido sensible.'],
['Canales y privacidad','HTTPS; avisos con contenido mínimo. Recopilar solo datos necesarios y acordar conservación y eliminación.'],
['Revisión de versiones','Revisar dependencias, migraciones y permisos antes de publicar; registrar hallazgos y correcciones.']],[3000,7600])
p('Usamos OWASP MASVS como referencia móvil para almacenamiento, autenticación y red [4]. Estos controles son una propuesta verificable, no una certificación de cumplimiento. En Compose no se aplican automáticamente pruebas XSS propias de HTML; priorizamos autorización por objeto, validación y protección de sesión según la superficie real.')
page('7 Plan de pruebas funcionales')
sub('7 1 Organización del ensayo')
p('Objetivo: comprobar que el pedido y la cotización atraviesan cliente, servidor y dueña sin perder integridad. Entorno: proyecto de ensayo con migraciones verificadas, APK identificado, dos dispositivos y tres cuentas ficticias (cliente A, cliente B y dueña). Registrar versión Android, red y permisos. No usar pedidos reales ni escribir sobre producción para la demostración.')
p('Responsables propuestos: Felipe ejecuta las pruebas de app e integración; Jafet registra resultados y revisa reglas y datos. Ambos revisan los resultados. Se confirma esta distribución antes del ensayo. Todos los casos siguientes quedan PLANIFICADOS; el presente informe no acredita su ejecución.')
table(['Caso y requisito','Precondición y pasos','Resultado esperado'],[
['PF-01 · RF-01/03/04/07','Producto de ensayo, fecha futura con 2 cupos. A solicita 1 unidad; dueña cotiza $25.000; A recarga y acepta; dueña entrega.','Mismo pedido en ambos dispositivos, cotización vigente, estado entregado y una sola venta.'],
['PF-02 · RF-02','Fecha de ensayo con 1 cupo y ninguna ocupación. A y B solicitan casi simultáneamente; cancelar el pedido aceptado.','Solo uno se guarda. El otro recibe rechazo. Cancelar libera cupo.'],
['PF-03 · RF-04','A abre cotización; dueña modifica monto antes de aceptación. A intenta confirmar su revisión anterior.','Rechazo de versión obsoleta; recarga permite revisar el nuevo monto.'],
['PF-04 · RF-08','Dueña crea producto con foto, edita opciones y retira después de asociar pedido.','Catálogo actualiza; pedido conserva referencia; retirado no se ofrece como activo.'],
['PF-05 · RF-05/06','Registrar venta y gasto de ensayo; editar monto y eliminar/anular según contrato.','Totales coinciden con registros vigentes; entregar dos veces no duplica venta.']],[2300,4450,3850])
p('Guardar evidencia por caso: ID, fecha, ejecutor, APK/commit, esperado, observado, captura anonimizada y defecto. La limpieza de datos se realiza solo en el entorno de ensayo, conservando los registros necesarios para explicar resultados.')
page('7 Seguridad y aceptación del diseño')
sub('7 2 Prueba de seguridad prioritaria')
p('PS-01 · RF-09 / RNF-02. Precondiciones: A y B tienen pedidos distintos; A es cliente sin rol de dueña. Pasos: con token de A intentar consultar o modificar el pedido de B mediante el contrato remoto, escribir un producto y consultar finanzas; repetir sin sesión; finalmente repetir operaciones autorizadas como dueña. Usar una herramienta API controlada y no registrar tokens en capturas.')
p('Resultado esperado: servidor niega modificaciones no autorizadas y no devuelve datos privados. El intento no altera el pedido de B, el catálogo ni las finanzas. El acceso legítimo de la dueña funciona. Comparar el estado antes/después desde cuenta autorizada. Un mensaje de error en la pantalla sin comprobación de datos no es evidencia suficiente.')
sub('7 3 Integración y experiencia')
table(['Caso','Procedimiento','Aceptación'],[
['PI-01 · RF-10','Con permiso y sesión recordada, poner app A en segundo plano; dueña cotiza; abrir aviso y luego cambiar de usuario.','Aviso pertinente abre pedidos; la nueva cuenta no recibe contenido de la anterior. Registrar restricciones del teléfono.'],
['PI-02 · RF-07','Interrumpir red antes de guardar; recuperar conexión y consultar antes de reintentar.','No anunciar éxito falso ni duplicar registros; informar estado recuperable.'],
['PU-01 · RNF-03','Dueña crea un producto y cliente solicita con texto ampliado; observar ayuda y errores.','Tareas completas sin bloqueo, opciones comprensibles y controles legibles. Registrar resultado real.']],[2300,4450,3850])
sub('7 4 Entrada salida y defectos')
p('Entrada: cuentas, datos, configuración y versión identificados; consentimiento para sesión de usabilidad. Salida propuesta: PF-01 y PS-01 aprobados como mínimo de guía, y ningún defecto crítico de acceso, pérdida de datos, sobrecupo o venta duplicada. Los demás casos amplían la cobertura del proyecto y cualquier pendiente debe quedar visible.')
p('Ante fallo: registrar pasos reproducibles, severidad, responsable, corrección y reprueba; repetir casos afectados. Las 39 pruebas unitarias históricas apoyan reglas locales, pero no sustituyen los ensayos remotos. En esta elaboración no se repitió Gradle ni una prueba en teléfonos.')
page('8 Conclusiones y cobertura U2')
p('El diseño vincula la necesidad de coordinación del negocio con una app de dos perfiles y un servidor que conserva reglas y datos comunes. La arquitectura separa interfaz, coordinación, dominio y acceso remoto. El modelo relacional evita referencias inconsistentes y la revisión de cotización protege decisiones tomadas con información antigua.')
p('La metodología incremental permite consolidar lo construido sin ampliar funciones por disponibilidad técnica. La propuesta de interfaz respeta la identidad del negocio y evita exigir opciones de pastel a otros productos. Las decisiones de seguridad se convierten en pruebas con resultados esperados observables.')
table(['Criterio','Evidencia del informe','Cierre pendiente'],[
['2.1.1 Cronograma','Sección 2: fechas propuestas, Gantt, 15 h, hitos, recursos y dependencias.','Acordar fecha oficial y responsables.'],
['2.1.2 Arquitectura 4+1','Sección 3: clases, componentes, BPMN, escenarios y despliegue.','Revisión técnica conjunta.'],
['2.1.3 Interfaces','Sección 4: perfiles, navegación, prototipo y criterios UX.','Validación observada con usuarios.'],
['2.1.4 Datos','Sección 5: entidades, atributos, relaciones, integración y límites.','Contrastar esquema remoto.'],
['2.1.5 Seguridad','Sección 6: privilegios, validaciones, sesión, errores y secretos.','Ensayo de permisos efectivos.'],
['2.1.6 Pruebas','Sección 7: funcional y seguridad, más integración/usabilidad.','Ejecutar y registrar en U3.']],[2600,5200,2800])
p('La profesora y sede se registran según confirmación del equipo. Sección/grupo y fecha de entrega siguen pendientes. Para entregar en AAI se utilizará el nombre Grupo XX_Sección XX.zip con los valores reales; no se inventa esa identificación. La regla de diapositivas se reservará para la presentación final.')
p('Esta versión es un informe de diseño listo para revisión del equipo. La correspondencia con los seis criterios facilita evaluar su cobertura, pero no asegura una nota ni reemplaza las indicaciones específicas de la docente.')
page('Referencias y fuentes')
p('[1] INACAP. Guía ABPro Unidad 2, Diseño de la solución de software, Proyecto Integrado TIHV43. Documento proporcionado por el equipo: TIHV43_GUÍA_ABPRO_U2.docx. Actividades 1–6 y criterios 2.1.1–2.1.6.')
p('[2] Alvarado, F. y Lavados, J. (2026). Diagnóstico del problema informático, Delicias de Mamá. Informe U1 del 7 de octubre de 2026, versión de revisión. Requisitos RF-01 a RF-10 y RNF-01 a RNF-06.')
p('[3] Equipo Delicias de Mamá. Código y documentación del proyecto Android, revisados el 13 de octubre de 2026. Repositorio: https://github.com/GelipeMlg/delicias-de-mama-android. Fuentes principales: Models.kt, MainActivity.kt, ShopViewModel.kt, ConnectedRepository.kt, SecureSessionStore.kt y migraciones SQL del 16 al 29 de septiembre de 2026.')
p('[4] OWASP Foundation. Mobile Application Security Verification Standard (MASVS). Consultado el 13 de octubre de 2026. https://mas.owasp.org/MASVS/ . Referencia para categorías de controles móviles; no certifica el proyecto.')
p('[5] Supabase. Row Level Security. Consultado el 13 de octubre de 2026. https://supabase.com/docs/guides/database/postgres/row-level-security . Referencia de autorización de datos; las funciones privilegiadas requieren análisis propio.')
sub('Anexos técnicos disponibles')
p('En Archivos Proyecto Integrado/03 Documentacion/U2/figuras se guardan los nueve diagramas y prototipo incluidos. Los scripts scripts/figuras_u2.py y scripts/generar_informe_u2_2026.py permiten regenerar los recursos y el Word. Los archivos SQL son la fuente detallada de tipos, restricciones y funciones; los diagramas seleccionan atributos para mantener legibilidad.')
sub('Registro de alcance de la revisión')
p('Se consultaron pauta, informe U1, código Android y migraciones. No se modificó la app ni la base remota al redactar U2. Las ilustraciones son modelos de diseño, no capturas ni resultados de prueba. Las recomendaciones de UX, la asignación de tareas y el cronograma esperan confirmación del equipo y de la contraparte.')
pp=prop(closing);setprop(pp,'pageBreakBefore');setprop(pp,'spacing',before=1440)
body.append(closing);body.append(sect)
parts['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
with ZipFile(REF) as a,ZipFile(OUT) as b:
 assert all(a.read(n)==b.read(n) for n in a.namelist() if n not in ('word/document.xml','word/_rels/document.xml.rels'))
 assert E.tostring(E.fromstring(a.read('word/document.xml')).find('w:body/w:sectPr',N))==E.tostring(E.fromstring(b.read('word/document.xml')).find('w:body/w:sectPr',N))
print('U2 creado; plantilla preservada, relaciones de figuras añadidas.')
